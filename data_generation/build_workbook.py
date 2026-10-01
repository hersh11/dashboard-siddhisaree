import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from gen_orders import build_dataset

employees, customers, orders, stages = build_dataset(1060)

df_emp = pd.DataFrame(employees)
df_cust = pd.DataFrame(customers)
df_orders = pd.DataFrame(orders)
df_stages = pd.DataFrame(stages)

# order_count / total_spend per customer for reference
spend = df_orders.groupby("customer_id")["final_price_inr"].agg(["count", "sum"]).reset_index()
spend.columns = ["customer_id", "total_orders", "total_spend_inr"]
df_cust = df_cust.merge(spend, on="customer_id", how="left")
df_cust["total_orders"] = df_cust["total_orders"].fillna(0).astype(int)
df_cust["total_spend_inr"] = df_cust["total_spend_inr"].fillna(0).astype(int)

HEADER_FILL = PatternFill("solid", start_color="C0392B", end_color="C0392B")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=10)

wb = Workbook()
wb.remove(wb.active)

def write_sheet(name, df, freeze="A2"):
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for c in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in df.itertuples(index=False):
        ws.append(list(row))
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).font = BODY_FONT
    for c, col in enumerate(df.columns, start=1):
        maxlen = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).values[:200]])
        ws.column_dimensions[get_column_letter(c)].width = min(max(maxlen + 2, 10), 55)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    return ws

write_sheet("Orders", df_orders)
write_sheet("OrderStages", df_stages)
write_sheet("Customers", df_cust)
write_sheet("Employees", df_emp)

# ---- README / Schema sheet ----
ws = wb.create_sheet("README", 0)
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 100
rows = [
    ("Dataset", "Siddhi Sarees Pvt. Ltd. — Bespoke Order Analytics Dataset (Synthetic)"),
    ("Period covered", "1 Jan 2026 – 30 Jun 2026 (as of 1 Jul 2026)"),
    ("Important note", "This is a SYNTHETIC dataset generated for prototyping the analytics tool. "
        "All names, notes and contact details are fictional / randomly generated. "
        "Replace with real exported data once available; keep the same column structure so your "
        "dashboards/queries continue to work unchanged."),
    ("", ""),
    ("Sheet: Orders", "One row per order (1,060 rows). Core order attributes, pricing, current stage/status."),
    ("Sheet: OrderStages", "One row per workflow stage per order (long/event format, ~7,880 rows). "
        "Tracks the 8-stage pipeline: Order Created -> Designer Assigned -> Design Approved -> "
        "Production Assigned -> Quality Check -> Packaging -> Shipping -> Delivered. "
        "Each row has assigned employee, start/end timestamp, status, notes and image references."),
    ("Sheet: Customers", "One row per unique customer (local Bareilly/UP, pan-India, and international NRI buyers), "
        "with total_orders / total_spend_inr computed from Orders."),
    ("Sheet: Employees", "Staff roster: Sales Executives, Designers, Production Karigars, Quality Checkers, "
        "Packaging Staff, Shipping Coordinators — reused consistently across orders."),
    ("Sheet: Summary", "Quick pivot-style KPIs computed live with formulas from the Orders sheet."),
    ("", ""),
    ("Order statuses", "Delivered / In Progress / On Hold / Cancelled"),
    ("Image references", "Filenames only (e.g. IMG_ORD1001_DesignApproved_1.jpg) as a placeholder for wherever "
        "actual stage photos will be stored (e.g. cloud storage / Drive folder linked by order_id + stage)."),
    ("Regenerating data", "This was generated with a Python script (gen_employees.py / gen_customers.py / "
        "gen_orders.py / build_workbook.py). Ask Claude to adjust volume, seasonality, price ranges, "
        "or add more product types and regenerate any time."),
]
for r in rows:
    ws.append(r)
ws["A1"].font = Font(bold=True, size=14)
for r in range(1, ws.max_row + 1):
    ws.cell(row=r, column=1).font = Font(bold=True)
    ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30

# ---- Summary sheet with live formulas ----
ws = wb.create_sheet("Summary")
ws.column_dimensions["A"].width = 32
ws.column_dimensions["B"].width = 18
n = len(df_orders) + 1  # +1 header row in Orders sheet

def col_letter(colname):
    return get_column_letter(list(df_orders.columns).index(colname) + 1)

status_col = col_letter("current_status")
price_col = col_letter("final_price_inr")
cat_col = col_letter("product_category")
returning_col = col_letter("is_returning_customer")
country_col = col_letter("customer_country")

summary_rows = [
    ("Metric", "Value"),
    ("Total Orders", f"=COUNTA(Orders!A2:A{n})"),
    ("Delivered Orders", f'=COUNTIF(Orders!{status_col}2:{status_col}{n},"Delivered")'),
    ("In Progress Orders", f'=COUNTIF(Orders!{status_col}2:{status_col}{n},"In Progress")'),
    ("On Hold Orders", f'=COUNTIF(Orders!{status_col}2:{status_col}{n},"On Hold")'),
    ("Cancelled Orders", f'=COUNTIF(Orders!{status_col}2:{status_col}{n},"Cancelled")'),
    ("Total Revenue (Delivered, INR)", f'=SUMIF(Orders!{status_col}2:{status_col}{n},"Delivered",Orders!{price_col}2:{price_col}{n})'),
    ("Average Order Value (INR)", f"=AVERAGE(Orders!{price_col}2:{price_col}{n})"),
    ("Orders from Returning Customers", f'=COUNTIF(Orders!{returning_col}2:{returning_col}{n},TRUE)'),
    ("Orders from New Customers", f'=COUNTIF(Orders!{returning_col}2:{returning_col}{n},FALSE)'),
    ("International (NRI) Orders", f'=COUNTIF(Orders!{country_col}2:{country_col}{n},"<>India")'),
    ("", ""),
    ("Orders by Category", ""),
    ("Saree", f'=COUNTIF(Orders!{cat_col}2:{cat_col}{n},"Saree")'),
    ("Lehenga", f'=COUNTIF(Orders!{cat_col}2:{cat_col}{n},"Lehenga")'),
    ("Suit/Anarkali", f'=COUNTIF(Orders!{cat_col}2:{cat_col}{n},"Suit/Anarkali")'),
    ("Gown", f'=COUNTIF(Orders!{cat_col}2:{cat_col}{n},"Gown")'),
    ("Blouse Only", f'=COUNTIF(Orders!{cat_col}2:{cat_col}{n},"Blouse Only")'),
]
for row in summary_rows:
    ws.append(row)
ws["A1"].font = HEADER_FONT
ws["B1"].font = HEADER_FONT
ws["A1"].fill = HEADER_FILL
ws["B1"].fill = HEADER_FILL
ws["A13"].font = Font(bold=True, italic=True)
for r in range(2, ws.max_row + 1):
    if r != 13:
        ws.cell(row=r, column=1).font = Font(bold=True)

from pathlib import Path
wb.save(Path(__file__).resolve().parent.parent / "siddhi_sarees_dataset.xlsx")
print("saved")
print(df_orders.shape, df_stages.shape, df_cust.shape, df_emp.shape)
